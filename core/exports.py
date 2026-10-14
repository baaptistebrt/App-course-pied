"""
Exports du plan d'un élève : Excel, CSV et document imprimable (HTML A4).
"""

from __future__ import annotations

import html
import io
from datetime import date, datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from .calculs import format_minutes, format_pct_range, format_race_time, fr_number, ratio_formula
from .config import DAY_LABELS, DAYS, DISPLAY_NAMES, DISTANCES, FREQ_LABELS, METHODS, PACE_LABELS, ZONE_COLORS
from .plan import is_race_day, pace_codes, week_summary

EXPORT_COLUMNS = ["semaine", "debut", "fin", "phase", "seance", "date", "km_env", "duree_env", "contenu",
                  "allure", "allure_txt", "km_spe", "temps_aff", "temps_spe_min", "frequence", "fait", "ressenti",
                  "commentaire", "remarque"]

MONTHS = ["janv.", "févr.", "mars", "avr.", "mai", "juin", "juil.", "août", "sept.", "oct.", "nov.", "déc."]
WEEKDAYS = ["lun.", "mar.", "mer.", "jeu.", "ven.", "sam.", "dim."]


def fr_date(d: date | None, year: bool = False, weekday: bool = False) -> str:
    if not isinstance(d, date):
        return ""
    txt = f"{d.day} {MONTHS[d.month - 1]}" + (f" {d.year}" if year else "")
    return f"{WEEKDAYS[d.weekday()]} {txt}" if weekday else txt


def pace_labels(value) -> str:
    return ", ".join(PACE_LABELS[c] for c in pace_codes(value))


def export_table(table: pd.DataFrame, done: set[str] | None = None, notes: dict | None = None) -> pd.DataFrame:
    """Plan à exporter ; `done` ajoute « Réalisée », `notes` ajoute « Ressenti » et « Commentaire »."""
    out = table.copy()
    ids = list(out.get("session_id", [""] * len(out)))
    out["fait"] = ["Oui" if done and sid in done else ("Non" if done is not None else "") for sid in ids]
    out["ressenti"] = [notes[sid].label if notes and sid in notes else "" for sid in ids]
    out["commentaire"] = [notes[sid].commentaire if notes and sid in notes else "" for sid in ids]
    out = out.reindex(columns=EXPORT_COLUMNS)
    out["allure"] = [pace_labels(c) for c in out["allure"]]
    out["frequence"] = [FREQ_LABELS.get(f, "") for f in out["frequence"]]
    out["temps_spe_min"] = pd.to_numeric(out["temps_spe_min"], errors="coerce").round(0)
    if done is None:
        out = out.drop(columns=["fait"])
    if notes is None:
        out = out.drop(columns=["ressenti", "commentaire"])
    return out.rename(columns=DISPLAY_NAMES)


def paces_table(sp, student) -> pd.DataFrame:
    """Tableau des allures retenues, avec la formule de la méthode utilisée."""
    rows = []
    for i, p in enumerate(sp.paces.values(), start=1):
        if sp.method == "vma":
            lo, hi = student.vma_pct[p.code]
            formula = f"{format_pct_range(lo, hi)} de VMA ({fr_number(student.vma)} km/h)"
        else:
            formula = ratio_formula(p.code, student.ratios)
        rows.append({"Zone": i, "Allure": p.label, "Correspond à": p.description, "Formule": formula,
                     "Vitesse": p.kmh_str, "Allure (min/km)": p.pace_str,
                     "Équivalence": f"{p.ref_label} en {format_minutes(p.ref_time_minutes)}" if p.ref_time_minutes else ""})
    return pd.DataFrame(rows)


def profile_rows(student, sp, frequency: int, race_time: str = "") -> list[tuple[str, str]]:
    dist = DISTANCES[student.distance]
    rows = [("Élève", student.display_name), ("E-mail", student.email), ("Course", student.race_name),
            ("Distance", dist.label), ("Méthode de calcul", METHODS[sp.method])]
    if student.objectif:
        rows.append(("Objectif", format_race_time(sp.race_seconds)))
    else:
        rows.append(("Temps estimé", format_race_time(sp.race_seconds)))
    if student.vma is not None:
        rows.append(("VMA", f"{fr_number(student.vma)} km/h"))
    rows += [("Date de course", student.date_course.strftime("%d/%m/%Y") if student.date_course else "non renseignée"),
             ("Séances par semaine", str(frequency)), ("Semaines de plan", str(student.plan.n_weeks))]
    if race_time:
        rows.append(("Temps réalisé", race_time))
    rows.append(("Généré le", datetime.now().strftime("%d/%m/%Y %H:%M")))
    return rows


def export_excel(student, table: pd.DataFrame, sp, frequency: int, done: set[str] | None = None,
                 race_time: str = "", notes: dict | None = None) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        export_table(table, done, notes).to_excel(writer, sheet_name="Plan", index=False)
        paces_table(sp, student).to_excel(writer, sheet_name="Allures", index=False)
        pd.DataFrame(profile_rows(student, sp, frequency, race_time), columns=["", "Valeur"]).to_excel(
            writer, sheet_name="Profil", index=False)
        header_fill, band = PatternFill("solid", fgColor="1B4F8C"), PatternFill("solid", fgColor="EAF0F7")
        for ws in writer.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            for col_cells in ws.columns:
                length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
                ws.column_dimensions[col_cells[0].column_letter].width = min(max(9, length + 2), 42)
        ws = writer.book["Plan"]
        headers = [c.value for c in ws[1]]
        date_cols = {headers.index(DISPLAY_NAMES[k]) + 1 for k in ("debut", "fin", "date")}
        shade, previous = False, None
        for row in ws.iter_rows(min_row=2):
            if row[0].value != previous:
                shade, previous = not shade, row[0].value
            for cell in row:
                if shade:
                    cell.fill = band
                if cell.column in date_cols:
                    cell.number_format = "DD/MM/YYYY"
    return buffer.getvalue()


def export_csv(table: pd.DataFrame, done: set[str] | None = None, notes: dict | None = None) -> bytes:
    """CSV pour Excel en français : séparateur « ; », virgule décimale, UTF-8 avec BOM."""
    df = export_table(table, done, notes)
    for col in (DISPLAY_NAMES["debut"], DISPLAY_NAMES["fin"], DISPLAY_NAMES["date"]):
        df[col] = [d.strftime("%d/%m/%Y") if isinstance(d, date) else "" for d in df[col]]
    return df.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig")


# ----------------------------------------------------------------------------
# Document imprimable
# ----------------------------------------------------------------------------
PRINT_CSS = """
@import url('https://fonts.googleapis.com/css2?family=Barlow:wght@400;500;600&family=Barlow+Condensed:wght@600;700;800&display=swap');
@page { size: A4; margin: 12mm 11mm; }
.doc { font-family: Barlow, Arial, sans-serif; color: #14233A; font-size: 10.5pt; line-height: 1.35; }
.doc * { -webkit-print-color-adjust: exact; print-color-adjust: exact; box-sizing: border-box; }
.doc h1 { font-family: 'Barlow Condensed', Arial Narrow, sans-serif; font-weight: 800; font-size: 26pt;
          margin: 0; line-height: 1; }
.doc .head { display: flex; justify-content: space-between; align-items: flex-end; gap: 16px;
             border-bottom: 3px solid #1B4F8C; padding-bottom: 8px; margin-bottom: 10px; }
.doc .goal { text-align: right; }
.doc .goal b { font-family: 'Barlow Condensed', sans-serif; font-size: 22pt; font-weight: 800; color: #1B4F8C; }
.doc .meta { color: #5B6B80; font-size: 9.5pt; }
.doc table { width: 100%; border-collapse: collapse; }
.doc th { text-align: left; font-weight: 600; font-size: 8.5pt; color: #5B6B80; border-bottom: 1px solid #C9D3DF;
          padding: 3px 5px; }
.doc td { padding: 4px 5px; border-bottom: 1px solid #E3E9F0; vertical-align: top; }
.doc .paces td { padding: 3px 5px; }
.doc .zone { display: inline-block; text-align: center; border-radius: 3px; font-weight: 700;
             padding: 0 4px; margin: 0 4px 2px 0; white-space: nowrap; }
.doc .week { break-inside: avoid; page-break-inside: avoid; margin-top: 12px; }
.doc .week h2 { font-family: 'Barlow Condensed', sans-serif; font-size: 13.5pt; margin: 0 0 2px; font-weight: 700; }
.doc .week h2 small { font-family: Barlow, sans-serif; font-size: 9pt; color: #5B6B80; font-weight: 500; }
.doc .phase { color: #1B4F8C; font-size: 9pt; margin-bottom: 3px; }
.doc .race td { font-weight: 600; background: #FFF4D1; }
.doc .num { white-space: nowrap; font-variant-numeric: tabular-nums; }
.doc .box { font-size: 13pt; line-height: 1; }
.doc .feel { color: #33465F; font-style: italic; font-size: 9pt; }
.doc .foot { margin-top: 14px; color: #5B6B80; font-size: 8.5pt; }
"""


def _zone_badges(value, paces) -> str:
    out = []
    for code in pace_codes(value):
        bg, fg = ZONE_COLORS[code]
        txt = paces[code].pace_str if code in paces else ""
        out.append(f"<span class='zone' style='background:{bg};color:{fg}'>{html.escape(PACE_LABELS[code])}</span>"
                   f"<span class='num'>{html.escape(txt)}</span> ")
    return "<br>".join(out)


def print_document(student, table: pd.DataFrame, sp, frequency: int, done: set[str] | None = None,
                   race_time: str = "", notes: dict | None = None) -> str:
    """Corps HTML du plan imprimable (sans <html>), utilisé par le bouton et par le téléchargement."""
    e = html.escape
    dist = DISTANCES[student.distance]
    race = f"Course le {fr_date(student.date_course, True)}" if student.date_course else "Date de course non renseignée"
    basis = (f"allures sur VMA {fr_number(student.vma)} km/h" if sp.method == "vma"
             else "allures sur l'objectif de temps")
    goal_label = "Objectif" if student.objectif else "Temps estimé"
    parts = [f"<div class='doc'><div class='head'><div><h1>{e(student.display_name)}</h1>"
             f"<div class='meta'>{e(student.race_name)}{f' ({e(dist.label)})' if student.course else ''} — {e(race)} — "
             f"{frequency} séances par semaine — {e(basis)}</div></div>"
             f"<div class='goal'><div class='meta'>{goal_label}</div><b>{e(format_race_time(sp.race_seconds))}</b>"
             + (f"<div class='meta'>Réalisé : <b style='font-size:13pt'>{e(race_time)}</b></div>" if race_time else "")
             + "</div></div>"]
    rows = "".join(
        f"<tr><td>{_zone_badges(p.code, {})}</td><td class='num'><b>{e(p.pace_str)} /km</b></td>"
        f"<td class='num'>{e(p.kmh_str)}</td><td>{e(p.description)}</td>"
        f"<td class='num'>{e(f'{p.ref_label} en {format_minutes(p.ref_time_minutes)}') if p.ref_time_minutes else ''}</td></tr>"
        for p in sp.paces.values())
    parts.append(f"<table class='paces'><tr><th>Allure</th><th>Rythme</th><th>Vitesse</th><th>Usage</th>"
                 f"<th>Équivalence</th></tr>{rows}</table>")
    if not table.empty:
        for position in sorted(table["position"].unique()):
            sub = table[table["position"] == position]
            first = sub.iloc[0]
            dates = (f" <small>du {fr_date(first['debut'])} au {fr_date(first['fin'])}</small>"
                     if isinstance(first["debut"], date) else "")
            summary = week_summary(sub)
            volume = ", ".join(x for x in (summary.km_str, summary.duration_str) if x)
            parts.append(f"<div class='week'><h2>Semaine {int(position) + 1}{dates}"
                         f"{f' <small>{e(volume)}</small>' if volume else ''}</h2>")
            if first["phase"]:
                parts.append(f"<div class='phase'>{e(first['phase'])}</div>")
            parts.append("<table><tr><th style='width:4%'></th><th style='width:13%'>Séance</th><th>Contenu</th>"
                         "<th style='width:6%'>Km</th><th style='width:8%'>Durée</th><th style='width:24%'>Allure</th>"
                         "<th style='width:8%'>Temps spé.</th><th style='width:15%'>Remarque</th></tr>")
            for _, r in sub.iterrows():
                race_row = is_race_day(r["seance"])
                checked = done is not None and (r["session_id"] in done or (race_row and race_time))
                when = fr_date(r.get("date"), weekday=True)
                label = e(str(r["seance"])) + (f"<br><span class='meta'>{e(when)}</span>" if when else "")
                content = e(str(r["contenu"])) + (f"<br><b>Réalisé : {e(race_time)}</b>" if race_row and race_time else "")
                note = (notes or {}).get(r["session_id"])
                if note is not None and not note.empty:
                    content += f"<br><span class='feel'>{e(note.label)}{(' : ' + e(note.commentaire)) if note.commentaire else ''}</span>"
                parts.append(
                    f"<tr{' class=race' if race_row else ''}><td class='box'>{'☑' if checked else '☐'}</td>"
                    f"<td>{label}</td><td>{content}</td>"
                    f"<td class='num'>{e(str(r['km_env']))}</td><td class='num'>{e(str(r['duree_env']))}</td>"
                    f"<td>{_zone_badges(r['allure'], sp.paces)}</td>"
                    f"<td class='num'>{e(str(r['temps_aff']))}</td><td>{e(str(r['remarque']))}</td></tr>")
            parts.append("</table></div>")
    parts.append(f"<div class='foot'>Plan généré le {datetime.now():%d/%m/%Y}. Allures calculées par ton coach "
                 f"({e(basis)}). Coche ☐ chaque séance réalisée.</div></div>")
    return "".join(parts)


def print_page(body: str, title: str) -> str:
    """Page HTML autonome (téléchargeable), qui s'imprime directement depuis le navigateur."""
    return (f"<!doctype html><html lang='fr'><head><meta charset='utf-8'><title>{html.escape(title)}</title>"
            f"<meta name='viewport' content='width=device-width, initial-scale=1'>"
            f"<style>{PRINT_CSS} body{{margin:24px;background:#fff}} @media print{{body{{margin:0}}}}</style>"
            f"</head><body>{body}</body></html>")


def print_button_html(body: str, label: str = "Imprimer le plan") -> str:
    """Bouton à intégrer dans l'application : imprime le document (masqué à l'écran)."""
    return f"""
<style>
{PRINT_CSS}
html, body {{ margin: 0; background: transparent; }}
.btn {{ font-family: Barlow, Arial, sans-serif; font-size: 15px; font-weight: 600; color: #fff; background: #1B4F8C;
        border: 1.5px solid #1B4F8C; border-radius: 8px; padding: 0 16px; height: 40px; width: 100%; cursor: pointer; }}
.btn:hover {{ background: #0F2A4A; }}
.btn:focus-visible {{ outline: 3px solid #FFC629; outline-offset: 2px; }}
@media screen {{ .doc {{ display: none; }} }}
@media print {{ .btn {{ display: none; }} html, body {{ background: #fff; }} }}
</style>
<button class="btn" onclick="window.print()">🖨️ {html.escape(label)}</button>
{body}
"""


__all__ = ["DAYS", "DAY_LABELS", "fr_date", "export_excel", "export_csv", "print_document", "print_page",
           "print_button_html", "paces_table"]
