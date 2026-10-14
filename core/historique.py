"""
État d'un programme (plan affiché, suivi, progression) et historique des courses terminées.

Un programme est « terminé » quand l'élève a saisi son temps de course, ou quand la date de la
course est passée. L'historique garde, pour chaque course terminée : le nom, la date, l'objectif,
le chrono, l'écart, la part du plan réalisée, les km courus, le ressenti et le récit de la course.
Rien n'est recopié : tout est relu dans les fichiers du programme et du suivi, donc toujours à jour.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from . import suivi as sv_mod
from .calculs import format_race_time, fr_number
from .config import DISTANCES, METHODS
from .moteur import StudentPaces, student_paces
from .plan import full_plan_table, is_race_day
from .stockage import Store


@dataclass
class ProgramState:
    student: object
    sp: StudentPaces
    frequency: int
    table: pd.DataFrame
    suivi: sv_mod.Suivi
    progress: sv_mod.Progress

    @property
    def status(self) -> dict[str, bool]:
        if self.table.empty:
            return {}
        return {r["session_id"]: sv_mod.is_done(r, self.suivi) for _, r in self.table.iterrows()}

    def finished(self, today: date | None = None) -> bool:
        return is_finished(self.student, self.suivi, today)


def program_state(store: Store, s) -> ProgramState:
    """Plan tel que l'élève le voit (sa fréquence) + son suivi + sa progression."""
    sp = student_paces(s)
    sv = sv_mod.load_suivi(store, s.id)
    freq = sv_mod.effective_frequency(s, sv)
    table = full_plan_table(s.plan, sp.paces, sp.race_seconds, s.date_course, freq, hide_empty=True)
    objective = sp.race_seconds if s.objectif else None
    return ProgramState(s, sp, freq, table, sv, sv_mod.progress(table, sv, objective))


def is_finished(s, suivi: sv_mod.Suivi, today: date | None = None) -> bool:
    today = today or date.today()
    return bool(suivi.temps_realise) or (s.date_course is not None and s.date_course < today)


# ----------------------------------------------------------------------------
# Historique
# ----------------------------------------------------------------------------
@dataclass(frozen=True)
class RaceRecord:
    student_id: str
    person: str
    name: str               # prénom de l'élève
    email: str
    course: str             # nom de la course (ou distance)
    distance: str
    race_date: date | None
    objective: str          # « 44'00 » (objectif ou temps estimé)
    objective_is_estimate: bool
    race_time: str          # « 43'10 » ou "" si non saisi
    race_seconds: int | None
    gap: str                # « −0'50 »
    gap_seconds: int | None
    done: int
    total: int
    km_done: float
    feeling: str            # « 😄 Excellent »
    comment: str
    method: str
    vma: float | None

    @property
    def completion(self) -> int:
        return round(100 * self.done / self.total) if self.total else 0


def race_record(state: ProgramState) -> RaceRecord:
    s, sv, pr = state.student, state.suivi, state.progress
    secs = None
    if sv.temps_realise:
        try:
            secs = sv_mod.parse_race_time(sv.temps_realise)
        except ValueError:
            secs = None
    race_rows = state.table[[is_race_day(x) for x in state.table["seance"]]] if not state.table.empty else []
    note = sv.note(race_rows["session_id"].iloc[0]) if len(race_rows) else sv_mod.Note()
    gap_secs = secs - state.sp.race_seconds if (secs is not None and s.objectif) else None
    return RaceRecord(
        student_id=s.id, person=s.person, name=s.display_name, email=s.email, course=s.race_name,
        distance=s.distance, race_date=s.date_course, objective=format_race_time(state.sp.race_seconds),
        objective_is_estimate=not s.objectif, race_time=pr.race_time, race_seconds=secs, gap=pr.race_gap,
        gap_seconds=gap_secs, done=pr.done, total=pr.total, km_done=pr.km_done, feeling=note.label,
        comment=note.commentaire, method=METHODS[state.sp.method], vma=s.vma)


def race_history(store: Store, students: list, today: date | None = None) -> list[RaceRecord]:
    """Courses terminées parmi `students`, de la plus récente à la plus ancienne."""
    records = []
    for s in students:
        state = program_state(store, s)
        if state.finished(today):
            records.append(race_record(state))
    return sorted(records, key=lambda r: (r.race_date or date.min, r.name), reverse=True)


def history_frame(records: list[RaceRecord], with_student: bool = True) -> pd.DataFrame:
    """Tableau lisible de l'historique (affichage et export)."""
    rows = []
    for r in records:
        row = {"Élève": r.name, "E-mail": r.email} if with_student else {}
        row.update({
            "Course": r.course, "Distance": DISTANCES[r.distance].label,
            "Date": r.race_date.strftime("%d/%m/%Y") if r.race_date else "",
            "Objectif": r.objective + (" (estimé)" if r.objective_is_estimate else ""),
            "Temps réalisé": r.race_time or "non saisi", "Écart": r.gap,
            "Plan réalisé": f"{r.completion} %", "Séances": f"{r.done}/{r.total}",
            "Km courus": fr_number(r.km_done), "Méthode": r.method,
            "VMA": fr_number(r.vma) if r.vma is not None else "",
            "Ressenti": r.feeling, "Récit de la course": r.comment,
        })
        rows.append(row)
    return pd.DataFrame(rows)


def history_excel(records: list[RaceRecord]) -> bytes:
    """Historique au format Excel : une ligne par course, temps en secondes pour les calculs."""
    df = history_frame(records)
    df.insert(df.columns.get_loc("Écart") + 1, "Écart (s)", [r.gap_seconds for r in records])
    df.insert(df.columns.get_loc("Temps réalisé") + 1, "Temps (s)", [r.race_seconds for r in records])
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Courses", index=False)
        ws = writer.book["Courses"]
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1B4F8C")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for col_cells in ws.columns:
            length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max(9, length + 2), 50)
    return buffer.getvalue()
